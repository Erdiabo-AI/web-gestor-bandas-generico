"""Generación de Excel (.xlsx) con openpyxl: ficha de banda."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .helpers import ahora_es

GRIS_OSCURO = PatternFill("solid", fgColor="3F3F46")   # cabecera de la web
BEIGE = PatternFill("solid", fgColor="E8E4BC")
ALTERNO = PatternFill("solid", fgColor="F4F4F5")
LINEA = Side(style="thin", color="999999")
BORDE = Border(left=LINEA, right=LINEA, top=LINEA, bottom=LINEA)
BLANCO_NEGRITA = Font(bold=True, color="FFFFFF")
AJUSTE = Alignment(wrap_text=True, vertical="center")
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _armamento(lista):
    if not isinstance(lista, list):
        return ""
    return ", ".join(f"{a.get('nombre')} ({a.get('precio')})"
                     for a in lista if isinstance(a, dict) and a.get("nombre"))


def _tabla_miembros(ws, fila, lineas, stats, estados):
    """Escribe la tabla de miembros desde `fila`; devuelve la siguiente fila libre."""
    cab = (["Nº", "Nombre", "Tipo", "Categoría", "Cant."]
           + [s.replace("_", " ").upper() for s in stats]
           + ["Armamento", "Exp.", "Coste unit.", "Coste total", "Estado", "Notas", "Observaciones"])
    for col, texto in enumerate(cab, 1):
        c = ws.cell(row=fila, column=col, value=texto)
        c.font, c.fill, c.border, c.alignment = BLANCO_NEGRITA, GRIS_OSCURO, BORDE, CENTRO
    centradas = {1, 5} | set(range(6, 6 + len(stats))) | {7 + len(stats)}   # Nº, Cant., stats, Exp.
    for i, ln in enumerate(lineas):
        fila += 1
        valores = ln["banda_lin_stats"] if isinstance(ln["banda_lin_stats"], dict) else {}
        datos = ([ln["banda_lin_orden"], ln["banda_lin_nombre"], ln["banda_lin_subtipo"],
                  ln["banda_lin_categoria"] or "", ln["banda_lin_cantidad"]]
                 + [valores.get(s, "") for s in stats]
                 + [_armamento(ln["banda_lin_armamento"]), ln["banda_lin_experiencia"] or 0,
                    ln["banda_lin_coste"], ln["banda_lin_coste"] * ln["banda_lin_cantidad"],
                    estados.get(ln["banda_lin_estado"] or 0, ""), ln["banda_lin_notas"] or "",
                    ln["banda_lin_observaciones"] or ""])
        for col, valor in enumerate(datos, 1):
            c = ws.cell(row=fila, column=col, value=valor)
            c.border = BORDE
            c.alignment = CENTRO if col in centradas else AJUSTE
            if i % 2:
                c.fill = ALTERNO
    return fila + 1, len(cab)


def build_banda_xlsx(b, plantilla, bajas, stats, estados_banda, estados_miembro, app_name):
    """Libro con una hoja «Ficha»: datos de la banda, miembros y bajas (sin historial)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Ficha"

    ws["A1"] = b["banda_nombre"]
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = (f"{b['banda_tipo']} · Campaña {b['campana_nombre']}"
                + (f" · {b['regla_descripcion']}" if b.get("regla_descripcion") else ""))
    ws["A2"].font = Font(color="555555")

    # Datos de la banda (fila 4 etiquetas, fila 5 valores)
    etiquetas = ["Usuario", "Banda", "Campaña", "Reglamento", "Tesorería", "Coste", "Estado"]
    valores = [b["usuario_nombre"] or "", b["banda_tipo"], b["campana_nombre"],
               b.get("regla_descripcion") or "", b["banda_tesoreria"] or 0,
               b["banda_valor_banda"] or 0, estados_banda.get(b["banda_estado"], b["banda_estado"])]
    for col, (e, v) in enumerate(zip(etiquetas, valores), 1):
        c = ws.cell(row=4, column=col, value=e)
        c.font, c.fill, c.border = Font(bold=True), BEIGE, BORDE
        c = ws.cell(row=5, column=col, value=v)
        c.border, c.alignment = BORDE, AJUSTE

    fila = 7
    if b["banda_observaciones"]:
        ws.cell(row=fila, column=1, value="Observaciones:").font = Font(bold=True)
        ws.cell(row=fila, column=2, value=b["banda_observaciones"])
        fila += 2

    ws.cell(row=fila, column=1, value=f"Miembros de la banda ({len(plantilla)})").font = Font(bold=True, size=12)
    fila, ncols = _tabla_miembros(ws, fila + 1, plantilla, stats, estados_miembro)

    if bajas:
        fila += 1
        ws.cell(row=fila, column=1, value=f"Bajas: despedidos y muertos ({len(bajas)})").font = Font(bold=True, size=12)
        fila, _ = _tabla_miembros(ws, fila + 1, bajas, stats, estados_miembro)

    ws.cell(row=fila + 1, column=1,
            value=f"{app_name} · generado el {ahora_es():%d/%m/%Y %H:%M}").font = Font(italic=True, size=8, color="777777")

    # Anchos de columna según el contenido (con tope) y ajustes de impresión
    for col in range(1, max(ncols, len(etiquetas)) + 1):
        letra = get_column_letter(col)
        largo = max((len(str(c.value)) for c in ws[letra] if c.value is not None and c.row > 3), default=6)
        ws.column_dimensions[letra].width = min(max(6, largo + 2), 40)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
